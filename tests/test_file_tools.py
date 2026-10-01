# -*- coding: utf-8 -*-
"""S5.3 文件能力测试：file_view 多格式读取 + file_write 写入（无需真实 Key）。

覆盖：
1. file_view：UTF-8 文本 + 按行分页 + 行数/截断 metadata
2. file_view：GBK 编码自动回退（中文环境刚需）
3. file_view：docx 真实解析（python-docx 生成 → 读回段落/表格）
4. file_view：xlsx 真实解析（openpyxl 生成 → 按 sheet 读回）
5. file_view：PDF 分派与容错（%PDF magic 识别；损坏 PDF 报错不抛异常）
6. file_write：overwrite / append / 目录自动创建
7. file_write：JSON 回读校验（合法通过；坏 JSON json_valid=False）
8. 权限：file_write on-demand 需审批、full-access 放行；file_view 两档放行
"""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

import sys

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from src.permissions import ApprovalOutcome, PermissionPolicy, SandboxMode
from src.tools.file_view import FileViewTool
from src.tools.file_write import FileWriteTool

FV = FileViewTool(max_bytes=65536, max_lines=200)
FW = FileWriteTool(max_bytes=65536)


def test_text_and_pagination():
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "note.txt"
        f.write_text("\n".join(f"line-{i}" for i in range(50)) + "\n", encoding="utf-8")
        # 完整读（默认 limit 200 → 全 50 行）
        r = FV.execute({"path": str(f)})
        assert r.ok and "line-49" in r.content and r.metadata["total_lines"] == 50, "应读完 50 行"
        # 分页：offset=10, limit=5（尾部带续读提示行）
        r = FV.execute({"path": str(f), "offset": 10, "limit": 5})
        assert r.ok and r.content.splitlines()[:5] == [f"line-{i}" for i in range(10, 15)], "分页应精确"
        assert "可设置 offset=15 继续" in r.content, "应提示续读位置"
        assert r.metadata["lines_shown"] == 5
        # 缺路径
        r = FV.execute({})
        assert not r.ok and "path" in r.error
        print("[1/8] file_view 文本 + 分页 通过")


def test_gbk_fallback():
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "gbk.txt"
        f.write_bytes("中文内容：端口 8000，配置文件".encode("gbk"))
        r = FV.execute({"path": str(f)})
        assert r.ok and "端口 8000" in r.content, "GBK 应回退解码成功（不乱码）"
        print("[2/8] file_view GBK 编码回退 通过")


def test_docx_read():
    try:
        import docx
    except ImportError:
        print("[3/8] 跳过（python-docx 未安装）")
        return
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "sample.docx"
        d = docx.Document()
        d.add_paragraph("第一段：项目介绍")
        d.add_paragraph("第二段：端口是 8000")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "键"
        t.cell(0, 1).text = "值"
        t.cell(1, 0).text = "端口"
        t.cell(1, 1).text = "8000"
        d.save(str(f))
        r = FV.execute({"path": str(f)})
        assert r.ok and r.metadata["format"] == "docx", "应识别为 docx"
        assert "端口是 8000" in r.content, "应提取段落文本"
        assert "8000" in r.content and "端口" in r.content, "应提取表格"
        print("[3/8] file_view docx 解析 通过")


def test_xlsx_read():
    try:
        import openpyxl
    except ImportError:
        print("[4/8] 跳过（openpyxl 未安装）")
        return
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "sample.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "模型"
        ws.append(["厂商", "模型名"])
        ws.append(["阿里云", "qwen3.8-max"])
        ws.append(["硅基流动", "deepseek-v4-flash"])
        wb.save(str(f))
        r = FV.execute({"path": str(f)})
        assert r.ok and r.metadata["format"] == "xlsx", "应识别为 xlsx"
        assert "模型" in r.content and "qwen3.8-max" in r.content, "应按 sheet 提取内容"
        print("[4/8] file_view xlsx 解析 通过")


def test_pdf_dispatch_and_fallback():
    with tempfile.TemporaryDirectory() as td:
        # ① magic=%PDF → 进入 PDF 分支；内容损坏 → 报错不抛异常
        bad = Path(td) / "broken.pdf"
        bad.write_bytes(b"%PDF-1.4\nthis is not a real pdf body")
        r = FV.execute({"path": str(bad)})
        assert not r.ok or "PDF" in r.content or "解析失败" in r.content, "损坏 PDF 应明确报错不抛异常"
        # ② 非 PDF 内容但 .pdf 扩展名 → magic 判定走文本
        fake = Path(td) / "fake.pdf"
        fake.write_text("这其实是个文本文件", encoding="utf-8")
        r = FV.execute({"path": str(fake)})
        assert r.ok and r.metadata["format"] == "text", "非 %PDF 开头应按文本处理"
        print("[5/8] file_view PDF 分派与容错 通过")


def test_write_overwrite_append():
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "out.txt"
        r = FW.execute({"path": str(f), "content": "第一行"})
        assert r.ok and f.read_text(encoding="utf-8") == "第一行", "overwrite 默认应写入"
        r = FW.execute({"path": str(f), "content": "第二行", "mode": "append"})
        assert r.ok and f.read_text(encoding="utf-8") == "第一行第二行", "append 应追加"
        # 覆盖：重写
        r = FW.execute({"path": str(f), "content": "覆盖内容"})
        assert f.read_text(encoding="utf-8") == "覆盖内容", "overwrite 应覆盖"
        # 新目录自动创建
        nested = Path(td) / "a" / "b" / "nested.json"
        r = FW.execute({"path": str(nested), "content": '{"ok": true}'})
        assert r.ok and nested.exists(), "父目录应自动创建"
        # 参数缺失
        r = FW.execute({"path": str(f)})
        assert not r.ok, "缺 content 应报错"
        r = FW.execute({"path": str(f), "content": "x", "mode": "bogus"})
        assert not r.ok, "非法 mode 应报错"
        print("[6/8] file_write 覆盖/追加/自动建目录 通过")


def test_write_json_validation():
    with tempfile.TemporaryDirectory() as td:
        f = Path(td) / "cfg.json"
        r = FW.execute({"path": str(f), "content": '{"port": 8000, "name": "agent"}'})
        assert r.ok and r.metadata["json_valid"] is True, "合法 JSON 应校验通过"
        bad = Path(td) / "bad.json"
        r = FW.execute({"path": str(bad), "content": '{"port": 8000, }'})
        assert r.ok and r.metadata["json_valid"] is False, "坏 JSON 应标记校验失败"
        # 非 JSON 内容不校验
        txt = Path(td) / "plain.txt"
        r = FW.execute({"path": str(txt), "content": "普通文本"})
        assert r.ok and r.metadata["json_valid"] is None, "非 JSON 不触发校验"
        print("[7/8] file_write JSON 回读校验 通过")


def test_permission_rules():
    p = PermissionPolicy()
    # file_write：on-demand → 审批；full-access → 放行
    out, _ = p.check(SandboxMode.ON_DEMAND.value, "file_write", {"path": "D:/x/out.txt", "content": "x"})
    assert out == ApprovalOutcome.NEEDS_APPROVAL, "on-demand 写文件应需审批"
    out, _ = p.check(SandboxMode.FULL_ACCESS.value, "file_write", {"path": "D:/x/out.txt", "content": "x"})
    assert out == ApprovalOutcome.ALLOWED, "full-access 写文件应放行"
    # file_view：两档放行
    out, _ = p.check(SandboxMode.ON_DEMAND.value, "file_view", {"path": "D:/x/in.txt"})
    assert out == ApprovalOutcome.ALLOWED, "file_view on-demand 应放行"
    out, _ = p.check(SandboxMode.FULL_ACCESS.value, "file_view", {"path": "D:/x/in.txt"})
    assert out == ApprovalOutcome.ALLOWED, "file_view full-access 应放行"
    print("[8/8] 权限规则（写审批/读放行） 通过")


def main():
    test_text_and_pagination()
    test_gbk_fallback()
    test_docx_read()
    test_xlsx_read()
    test_pdf_dispatch_and_fallback()
    test_write_overwrite_append()
    test_write_json_validation()
    test_permission_rules()
    print("\n✅ 全部通过：文件能力（多格式读取 + 安全写入 + 权限）可用")


if __name__ == "__main__":
    main()
